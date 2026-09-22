"""
NBFC ingestion — the third pipeline, and the case that tests whether the architecture
generalises or merely claims to.

The point of these tests is NOT that NBFC works. It is that it works through the SHARED
machinery: a manifest entry, a reader and a consolidator, and nothing else. If a future
change makes this source need its own compute module, its own column names, or its own
denominator rule, one of these fails.
"""
import json
import sys
from pathlib import Path

import pytest

ANALYSIS = next(p for p in Path(__file__).resolve().parents if (p / ".git").is_dir()) / "analysis"
sys.path.insert(0, str(ANALYSIS))

from core import manifest                                              # noqa: E402
from signals.compute import csv_sector                                 # noqa: E402
from pipelines.nbfc import extract_nbfc, consolidate_nbfc              # noqa: E402

ROOT = ANALYSIS.parent
DATA = ANALYSIS / "rbi_nbfc"


# ── the reader ────────────────────────────────────────────────────────────────

def test_a_release_carries_five_dated_columns_not_one():
    """The fact that turns three files into eleven dates, and the reason all eleven are
    registered from day one rather than backfilled in a later session."""
    for sections in sorted(DATA.glob("*/sections.json")):
        doc = json.loads(sections.read_text())
        assert len(doc["dates"]) == 5, f"{doc['source_file']} carries {len(doc['dates'])}"
        assert doc["period"] == max(doc["dates"]), "the period IS the latest dated column"


def test_the_hierarchy_is_read_from_the_codes_and_every_child_resolves_a_parent():
    doc = json.loads((DATA / "2026-07-31" / "sections.json").read_text())
    codes = {s["code"] for s in doc["sectors"]}
    assert "T" in codes, "no total ⇒ every share loses its denominator"
    for s in doc["sectors"]:
        if s["code"] == "T":
            assert s["level"] == 0 and s["parent_code"] == ""
            continue
        assert s["level"] == s["code"].count(".") + 1
        assert s["parent_code"] in codes, f"{s['code']} hangs off nothing"


def test_labels_lose_rbis_numbering_footnotes_and_of_which():
    doc = json.loads((DATA / "2026-07-31" / "sections.json").read_text())
    names = {s["code"]: s["sector"] for s in doc["sectors"]}
    assert names["4.1"] == "Housing Loans", "the footnote asterisk must not enter the name"
    assert names["2"] == "Industry", "'of which' is a note about the ROWS, not part of the name"
    for n in names.values():
        assert "of which" not in n.lower() and "*" not in n
        assert not n[0].isdigit(), f"{n!r} still carries RBI's numbering"


def test_a_reindented_row_stops_the_ingest(tmp_path):
    """Loud, not tolerant. A row silently re-indented becomes a child of the wrong parent,
    and every share under it would be measured against the wrong total."""
    import openpyxl
    src = next(DATA.glob("*/raw/*.xlsx"))
    wb = openpyxl.load_workbook(src)
    ws = wb["Press Release"]
    for r in range(7, 22):                      # move "2.1 Infrastructure" one column left
        if str(ws.cell(r, 3).value or "").strip().startswith("2.1 "):
            ws.cell(r, 2).value = ws.cell(r, 3).value
            ws.cell(r, 3).value = None
            break
    bad = tmp_path / "reindented.xlsx"
    wb.save(bad)
    with pytest.raises(SystemExit, match="re-indented"):
        extract_nbfc.parse(bad)


# ── the consolidated store ────────────────────────────────────────────────────

def test_three_releases_become_eleven_dates():
    docs = consolidate_nbfc.load_releases()
    rows, _ = consolidate_nbfc.consolidate(docs)
    dates = {r["date"] for r in rows}
    assert len(docs) == 3 and len(dates) == 11
    assert len(rows) == 15 * 11, "fifteen sectors on every date, or a sector went missing"


def test_a_later_release_wins_and_the_disagreement_is_reported():
    """RBI restates. Last release wins — the same single-source-of-truth rule SIBC applies —
    but a revision nobody saw is indistinguishable from a parsing bug, so it is printed."""
    docs = consolidate_nbfc.load_releases()
    rows, revisions = consolidate_nbfc.consolidate(docs)
    by_key = {(r["date"], r["code"]): r for r in rows}
    for doc in docs:                            # the newest release owns every date it names
        for s in doc["sectors"]:
            for d, v in s["values"].items():
                latest = max(x["period"] for x in docs if d in x["dates"])
                if doc["period"] == latest:
                    assert by_key[(d, s["code"])]["outstanding_cr"] == v
    assert isinstance(revisions, list)          # reported, never swallowed


# ── the contract with the shared compute module ───────────────────────────────

def test_the_csv_honours_the_shared_column_contract():
    """These names are not NBFC's to choose. `csv_sector` verifies them on load and means
    the same thing by them in every source — which is why this pipeline needs no module."""
    df = csv_sector._load_df("nbfc")
    for col in csv_sector.REQUIRED_COLUMNS:
        assert col in df.columns, f"{col} missing — csv_sector would refuse to load this"
    assert len(df) == 165


def test_nbfc_declares_the_shared_compute_shape_and_no_optional_columns():
    man = manifest.load("nbfc")
    assert man["compute_module"] == "csv_sector"
    schema = man["schema"]
    assert "scope_column" not in schema, "this source publishes one statement"
    assert "memo_flag_column" not in schema, "and no memo lens"


def test_the_shared_methods_compute_against_nbfc_with_no_nbfc_specific_code():
    df = csv_sector._load_df("nbfc")
    main = {"parent_code": "T", "child_level": 1, "entity_type": "sector", "window": 12}
    yoy = csv_sector.csv_sector_scan_yoy(main, "2026-07-31", df)
    assert len(yoy) == 5, "one row per main sector"

    size = csv_sector.csv_sector_scan_abs(main, "2026-07-31", df)
    parts = [r for r in size if r["entity_type"] == "sector"]
    agg = [r for r in size if r["entity_type"] == "aggregate"]
    assert len(parts) == 5 and len(agg) == 1, "the size scan adds the cut's own total"
    # And on THIS cut that total is the published one, because the parts are the whole.
    assert abs(agg[0]["value"] - sum(r["value"] for r in parts)) < 0.01


# ── the denominator rule, on the source that motivated it ─────────────────────

def test_the_main_cut_is_additive_and_the_four_of_which_cuts_are_not():
    """The reason phase 3 came before phase 4. Four of five NBFC cuts are "of which"
    decompositions, so on this source the inflated share is the COMMON case, not the
    exception — Services would have published CRE at 37.2% instead of 18.6%."""
    df = csv_sector._load_df("nbfc")
    def cut(parent, level):
        c = {"parent_code": parent, "child_level": level, "entity_type": "sector",
             "window": 12, "coherence_min": 0.9}
        rows = csv_sector.csv_sector_allocation(c, "2026-07-31", df)
        alloc = [r["value"] for r in rows if r["entity_type"] == "alloc"]
        cov = [r["value"] for r in rows if r["entity_type"] == "coverage"]
        return sum(alloc), (cov[0] if cov else None)

    total, coverage = cut("T", 1)
    assert coverage is None, "the five main sectors sum to the total to the rupee"
    assert abs(total - 100.0) < 0.01

    for parent, level, ceiling in (("2", 2, 95.0), ("2.1", 3, 80.0),
                                   ("3", 2, 60.0), ("4", 2, 80.0)):
        total, coverage = cut(parent, level)
        assert coverage is not None, f"{parent} is an 'of which' cut and must say so"
        assert 0 < coverage < 100
        assert total < ceiling, f"{parent} shares sum to {total:.1f} — divided by the parts?"


def test_the_published_yoy_check_actually_covers_the_store():
    """A gate that checks nothing passes too. This one must be examining real windows."""
    from pipelines.nbfc import validate_published_yoy as v
    import io, contextlib
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        assert v.main() == 0
    out = buf.getvalue()
    n = int(out.split("all ")[1].split(" values")[0])
    assert n >= 80, f"only {n} YoY values checked against RBI's own figure"


# ── the declared "of which" cuts must match what the arithmetic finds ─────────

def test_the_profiles_non_exhaustive_declaration_matches_what_the_compute_measures():
    """The skeleton profile DECLARES which cuts are 'of which' so the structural additivity
    check reports an expected residual instead of four permanent warnings. That is a second
    description of a fact the denominator rule measures every period — so it is guarded here
    rather than left free to drift, which is the same 'declared == computed' shape the
    card-cut contract uses."""
    profile = json.loads((DATA / "skeleton_profile.json").read_text())
    declared = {e["code"] for e in profile["exhaustive_decompositions"]["non_exhaustive"]}

    df = csv_sector._load_df("nbfc")
    period = max(df["date"].astype(str))
    measured = set()
    for parent, level in (("T", 1), ("2", 2), ("2.1", 3), ("3", 2), ("4", 2)):
        c = {"parent_code": parent, "child_level": level, "entity_type": "sector", "window": 12}
        rows = csv_sector.csv_sector_allocation(c, period, df)
        if any(r["entity_type"] == "coverage" for r in rows):
            measured.add(parent)
    assert declared == measured, (
        f"profile declares {sorted(declared)} non-exhaustive, the arithmetic finds "
        f"{sorted(measured)} — the declaration has drifted from the data")


def test_the_model_declares_why_it_has_no_behavioral_layer():
    """A source can legitimately have no causal layer yet. But a deferral and a forgotten
    layer look identical from the validator, so the reason is declared and the validator
    accepts only a declared one."""
    model = json.loads((DATA / "merged" / "system_model.json").read_text())
    meta = model["_meta"]
    assert not model.get("force_instances"), "if forces now exist, drop the deferral"
    assert meta.get("behavioral_status") == "deferred"
    assert len(meta.get("behavioral_deferred_reason", "")) > 60, "a reason, not a shrug"


def test_every_nbfc_dimension_shows_a_band_and_the_one_part_cuts_say_why():
    """Five dimensions, five blocks, every one speaking. The two cuts RBI gives a single
    child cannot have a mix — a mix needs something to be a mix BETWEEN — and a row that
    simply vanished would read as broken (DASHBOARD_SPEC §16)."""
    band = json.loads((ROOT / "web/public/data/nbfc_state.json").read_text())["dimensions"]
    assert set(band) == {"mainSectors", "industry", "infrastructure", "services", "retail"}
    for dim, blocks in band.items():
        for b in blocks:
            assert b["speed"], f"{dim} has no speed line"
            assert b["mix"] or b["no_mix_note"], f"{dim} is silent about its mix"
    for dim in ("industry", "infrastructure"):
        b = band[dim][0]
        assert b["mix"] is None and "only one part" in b["no_mix_note"]
