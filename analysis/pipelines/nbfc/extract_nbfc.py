#!/usr/bin/env python3
"""
extract_nbfc.py — one RBI NBFC release → {period}/sections.json

RBI's "Sectoral Deployment of Outstanding Credit by NBFCs (including HFCs)" is a single
sheet of fifteen sector rows in rupees crore. It is `csv_sector`-shaped — one measure over a
code hierarchy — so it needs no compute module of its own, only this reader.

TWO THINGS ARE UNLIKE SIBC, AND BOTH SIMPLIFY THE PIPELINE:

  * **A release carries FIVE dated columns, not one** — M-24, the prior-prior FY end, M-12,
    the prior FY end, and M. Three releases therefore hold ELEVEN distinct dates, seven of
    which can compute a YoY immediately. All of them are extracted; the consolidator decides
    what to do when two releases carry the same date.

  * **The column header IS the data date.** There is nothing to remap. SIBC needs a date
    gate because RBI publishes its Statement 1 on a Friday that can fall in the following
    month; here the period simply IS the month-end in the header, so this pipeline has no
    stage 1a and cannot misdate a period.

The release also prints its OWN year-on-year columns. They are extracted and kept so
`validate_published_yoy.py` can check our arithmetic against the source's — an external
check on our own computation that neither existing pipeline has.

    python3 analysis/pipelines/nbfc/extract_nbfc.py {xlsx}
    python3 analysis/pipelines/nbfc/extract_nbfc.py {xlsx} --print-period
"""
from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

ROOT = next(p for p in Path(__file__).resolve().parents if (p / ".git").is_dir())
sys.path.insert(0, str(ROOT / "analysis"))

DATA_DIR = ROOT / "analysis" / "rbi_nbfc"
SHEET = "Press Release"

#: The total's code. RBI numbers its sectors 1..5 and leaves the total unnumbered; every
#: level-1 sector hangs off this so the hierarchy has one root, as `csv_sector` expects.
TOTAL_CODE = "T"
TOTAL_LABEL = "NBFC Credit"

#: Which spreadsheet column an indent level is written in. The label's own CODE also implies
#: its depth ("2.1.1" is three levels down), and the two are cross-checked: a disagreement
#: means RBI re-indented the statement, which is a format change worth stopping for.
INDENT_COL = {2: 1, 3: 2, 4: 3}      # column B/C/D → level 1/2/3

_CODE = re.compile(r"^\s*(\d+(?:\.\d+)*)\.?\s+(.*)$", re.S)


def clean_label(raw: str) -> str:
    """The sector's name, without RBI's numbering, footnote markers or "of which" tail.

    "of which" is dropped rather than recorded: whether a cut's parts are the whole of it is
    MEASURED from the numbers every period (signals/README.md, the denominator rule), never
    read off a label. A label can go stale; the arithmetic cannot.
    """
    txt = re.sub(r"\bof\s+which\b", " ", raw, flags=re.I)
    txt = txt.replace("\n", " ").replace("*", " ")
    return re.sub(r"\s+", " ", txt).strip(" .")


def parse(xlsx: Path) -> dict:
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    if SHEET not in wb.sheetnames:
        raise SystemExit(f"{xlsx.name}: no '{SHEET}' sheet — found {wb.sheetnames}")
    ws = wb[SHEET]

    # ── the dated columns ─────────────────────────────────────────────────────
    dates: dict[int, str] = {}
    header_row = None
    for r in range(1, 12):
        found = {c: ws.cell(r, c).value for c in range(1, ws.max_column + 1)
                 if isinstance(ws.cell(r, c).value, datetime)}
        if len(found) >= 3:
            header_row = r
            dates = {c: v.strftime("%Y-%m-%d") for c, v in found.items()}
            break
    if not dates:
        raise SystemExit(f"{xlsx.name}: no row of dates found — the header moved, or the "
                         f"dates are no longer stored as dates")

    # ── the published YoY columns, keyed by the date they describe ────────────
    # "Jul 2026 / Jul 2025" describes the LATEST date; "Jul 2025 / Jul 2024" the middle one.
    published_yoy: dict[int, str] = {}
    for c in range(max(dates) + 1, ws.max_column + 1):
        label = ws.cell(header_row, c).value
        if not isinstance(label, str) or "/" not in label:
            continue
        newer = label.split("/")[0].strip()
        for d in dates.values():
            if datetime.strptime(d, "%Y-%m-%d").strftime("%b %Y") == newer:
                published_yoy[c] = d
                break

    # ── the sector rows ───────────────────────────────────────────────────────
    rows, seen = [], set()
    for r in range(header_row + 1, ws.max_row + 1):
        label_col = next((c for c in (2, 3, 4) if ws.cell(r, c).value), None)
        if label_col is None:
            continue
        raw = str(ws.cell(r, label_col).value)
        if raw.strip().lower().startswith(("notes", "source", "*")):
            break

        m = _CODE.match(raw)
        if m is None:
            if clean_label(raw).lower() != TOTAL_LABEL.lower():
                continue                        # a stray note line, not a sector
            code, name, level, parent = TOTAL_CODE, TOTAL_LABEL, 0, ""
        else:
            code, name = m.group(1), clean_label(m.group(2))
            level = code.count(".") + 1
            if INDENT_COL.get(label_col) != level:
                raise SystemExit(
                    f"{xlsx.name} row {r}: code {code!r} implies level {level} but it is "
                    f"written in column {label_col} (level {INDENT_COL.get(label_col)}). "
                    f"RBI re-indented the statement — check the format before ingesting.")
            parent = code.rsplit(".", 1)[0] if "." in code else TOTAL_CODE

        if code in seen:
            raise SystemExit(f"{xlsx.name}: sector code {code!r} appears twice")
        seen.add(code)

        values = {d: ws.cell(r, c).value for c, d in dates.items()
                  if isinstance(ws.cell(r, c).value, (int, float))}
        rows.append({
            "code": code, "sector": name, "level": level, "parent_code": parent,
            "values": values,
            "published_yoy": {d: ws.cell(r, c).value for c, d in published_yoy.items()
                              if isinstance(ws.cell(r, c).value, (int, float))},
        })

    if not rows:
        raise SystemExit(f"{xlsx.name}: no sector rows parsed")
    if not any(r["code"] == TOTAL_CODE for r in rows):
        raise SystemExit(f"{xlsx.name}: no '{TOTAL_LABEL}' total row — the statement's shape "
                         f"changed, and every share would lose its denominator")

    return {
        "source_file": xlsx.name,
        "statement": str(ws.cell(2, 2).value or "").strip(),
        # The PERIOD is the latest dated column — the data month. No remapping exists or is
        # needed: unlike SIBC's fortnightly Friday releases, the header states the date.
        "period": max(dates.values()),
        "dates": sorted(dates.values()),
        "sectors": rows,
    }


def resolve_period(xlsx: Path) -> str:
    return parse(xlsx)["period"]


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("xlsx", type=Path)
    ap.add_argument("--print-period", action="store_true",
                    help="resolve the period from the file and print it (gate contract)")
    a = ap.parse_args()

    if a.print_period:
        print(resolve_period(a.xlsx))
        return 0

    doc = parse(a.xlsx)
    out_dir = DATA_DIR / doc["period"]
    (out_dir / "raw").mkdir(parents=True, exist_ok=True)
    # Archive the source beside its output. SIBC's extractor did NOT do this while its
    # consolidator globbed only those raw dirs, so a genuinely new month could be ingested,
    # report success, and add nothing at all.
    archived = out_dir / "raw" / a.xlsx.name
    if a.xlsx.resolve() != archived.resolve():
        shutil.copy2(a.xlsx, archived)

    (out_dir / "sections.json").write_text(json.dumps(doc, indent=2) + "\n")
    print(f"✓ {doc['period']}  ·  {len(doc['sectors'])} sectors  ·  "
          f"{len(doc['dates'])} dated columns ({doc['dates'][0]} … {doc['dates'][-1]})")
    print(f"  → {(out_dir / 'sections.json').relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
